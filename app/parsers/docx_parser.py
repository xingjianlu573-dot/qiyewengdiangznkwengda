# -*- coding: utf-8 -*-
"""Word（.docx）解析：使用 python-docx，识别标题样式并按段落分组。"""
from typing import List, Dict

from docx import Document


# 常见标题样式名（Word 内置样式）
_HEADING_STYLES = {
    "heading 1", "heading 2", "heading 3", "heading 4",
    "标题 1", "标题 2", "标题 3", "标题 4",
    "title", "标题",
}


def parse_docx(file_path: str, doc_name: str) -> List[Dict]:
    """解析 .docx，按标题样式切分章节；无标题的段落归入当前章节。"""
    try:
        document = Document(file_path)
    except Exception as exc:
        # 旧版 .doc（OLE 复合文档）python-docx 无法打开：给出可操作的提示
        if "Package not found" in str(exc) or "not a zip" in str(exc).lower():
            raise ValueError(
                f"无法打开《{doc_name}》：旧版 .doc 格式不支持，请在 Word 中另存为 .docx 后再上传"
            ) from exc
        raise ValueError(f"无法打开 Word 文档：{exc}") from exc

    blocks: List[Dict] = []
    current_heading = doc_name
    current_texts: List[str] = []
    order = 0

    def flush():
        nonlocal current_texts
        if current_texts:
            body = "\n".join(t.strip() for t in current_texts if t.strip())
            if body:
                blocks.append({
                    "heading": current_heading,
                    "text": body,
                    "page": None,
                    "order": order,
                })
        current_texts = []

    for para in document.paragraphs:
        style_name = (para.style.name or "").strip().lower()
        text = para.text.strip()
        if not text:
            continue
        if style_name in _HEADING_STYLES:
            flush()
            current_heading = text
            order += 1
        else:
            current_texts.append(text)

    flush()

    # 补充表格内容（很多故障处理文档用表格列步骤）
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append({
                "heading": current_heading,
                "text": "\n".join(rows),
                "page": None,
                "order": order + 1,
            })

    if not blocks:
        raise ValueError("Word 文档中未提取到任何文本内容")
    return blocks
